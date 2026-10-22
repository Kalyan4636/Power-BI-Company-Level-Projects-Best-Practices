"""Build the formula-driven Excel sales dashboard.

Usage:
    python src/build_dashboard.py --data data/sales_data.csv --out dashboard/Sales_Performance_Dashboard.xlsx
    python src/build_dashboard.py --image-mode --out build/dashboard_image.xlsx   # variant used for the PNG export

The workbook is generated with openpyxl. Every number on the dashboard is a live Excel formula;
no values are pasted in. See docs/architecture.md for the design.
"""
import argparse, csv, os, re, shutil, warnings, zipfile
from datetime import datetime
import openpyxl
warnings.filterwarnings('ignore')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, LineChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.series import DataPoint
from openpyxl.utils import get_column_letter as L


_p=argparse.ArgumentParser(description=__doc__.split("\n")[0])
_p.add_argument('--data',default=os.path.join(os.path.dirname(os.path.abspath(__file__)),'..','data','sales_data.csv'))
_p.add_argument('--out',default='dashboard/Sales_Performance_Dashboard.xlsx')
_p.add_argument('--image-mode',action='store_true',help='drop nav buttons / clutter for a static image export')
ARGS=_p.parse_args()
IMG=ARGS.image_mode
OUT=ARGS.out
os.makedirs(os.path.dirname(os.path.abspath(OUT)),exist_ok=True)
def load_rows(path):
    """Read the 12 source columns from CSV and restore types (ID stays text, dates become datetime)."""
    out=[]
    with open(path,newline='',encoding='utf-8') as f:
        rd=csv.reader(f); next(rd)
        for r in rd:
            out.append([r[0],datetime.strptime(r[1],'%Y-%m-%d'),r[2],r[3],r[4],r[5],r[6],r[7],
                        int(r[8]),float(r[9]),float(r[10]),float(r[11])])
    return out
rows=load_rows(ARGS.data); N=len(rows)

# palette
NAVY='1F2A44'; TEAL='0F8B8D'; CORAL='EE6C4D'; SAND='F4F1EA'; GREY='6B7280'; WHITE='FFFFFF'; LIGHT='E8F1F2'
F='Arial'
def font(sz=10,b=False,c='000000',i=False): return Font(name=F,size=sz,bold=b,color=c,italic=i)
def fill(c): return PatternFill('solid',fgColor=c)
thin=Side(style='thin',color='D1D5DB')


from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Protection
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font as DFont
from openpyxl.worksheet.properties import PageSetupProperties

wb=Workbook()
dash=wb.active; dash.title='Dashboard'
ana=wb.create_sheet('Analysis'); dat=wb.create_sheet('Data'); lst=wb.create_sheet('Lists'); gd=wb.create_sheet('Guide')

def nav(ws,row,cols,cur):
    items=[('Dashboard','Dashboard'),('Analysis','Analysis'),('Data','Data'),('Guide','Guide')]
    for (c1,c2),(label,sh) in zip(cols,items):
        cell=ws.cell(row,c1,label)
        if c2>c1: ws.merge_cells(start_row=row,start_column=c1,end_row=row,end_column=c2)
        for cc in range(c1,c2+1): ws.cell(row,cc).fill=fill(CORAL if sh==cur else TEAL)
        cell.alignment=Alignment(horizontal='center',vertical='center')
        cell.hyperlink=Hyperlink(ref=cell.coordinate,location=f"'{sh}'!A1",display=label)
        cell.font=font(10,True,WHITE)

# ---------- Lists
cats=sorted({r[6] for r in rows}); chans=sorted({r[4] for r in rows}); custs=sorted({r[5] for r in rows})
regions=sorted({r[2] for r in rows}); countries=sorted({r[3] for r in rows}); prods=sorted({r[7] for r in rows})
months=['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
lst['A1']='Lists (feed the dashboard dropdowns - do not edit)'; lst['A1'].font=font(11,True)
def put_list(col,title,items,allopt=True):
    lst.cell(3,col,title).font=font(10,True,WHITE); lst.cell(3,col).fill=fill(NAVY)
    vals=(['All'] if allopt else [])+items
    for i,v in enumerate(vals): lst.cell(4+i,col,v).font=font(10)
    lst.column_dimensions[L(col)].width=22
    return f"Lists!${L(col)}$4:${L(col)}${3+len(vals)}"
LR=put_list(1,'Region',regions); LCh=put_list(2,'Sales Channel',chans); LCu=put_list(3,'Customer Type',custs)
LM=put_list(4,'Month',months,False); LCa=put_list(5,'Category',cats); LCo=put_list(6,'Country',countries); LP=put_list(7,'Product',prods)
lst.sheet_properties.tabColor='9CA3AF'

# ---------- Data
dat['A1']='Sales Data'; dat['A1'].font=font(14,True,NAVY)
nav(dat,1,[(8,9),(10,11),(12,13),(14,15)],'Data')
hdr=['ID','Order Date','Region','Country','Sales Channel','Customer Type','Product Category','Product','Units Sold','Unit Price','Discount %','Unit Cost','Revenue','COGS','Gross Profit','Gross Margin %','Year','Quarter','MonthNo','Month','In Filter',
     'Flag excl. Country','Flag excl. Channel','Flag excl. Customer','Flag excl. Category','Flag excl. Product','Flag excl. Month']
HR=3; R0=4; R1=HR+N; NC=len(hdr)
for j,h in enumerate(hdr,1):
    c=dat.cell(HR,j,h); c.font=font(10,True,WHITE); c.fill=fill(NAVY if j<=21 else GREY); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
def flag(x,skip=None):
    cond={'R':f'OR(sel_Region="All",C{x}=sel_Region)','Co':f'OR(sel_Country="All",D{x}=sel_Country)','Ch':f'OR(sel_Channel="All",E{x}=sel_Channel)',
          'Cu':f'OR(sel_Customer="All",F{x}=sel_Customer)','Ca':f'OR(sel_Category="All",G{x}=sel_Category)','Pr':f'OR(sel_Product="All",H{x}=sel_Product)',
          'M':f'AND(S{x}>=m_From,S{x}<=m_To)'}
    return '=IF(AND('+','.join(v for k,v in cond.items() if k!=skip)+'),1,0)'
for i,r in enumerate(rows):
    x=R0+i
    for j,v in enumerate(r,1): dat.cell(x,j,v)
    dat.cell(x,13,f'=I{x}*J{x}*(1-K{x})'); dat.cell(x,14,f'=I{x}*L{x}'); dat.cell(x,15,f'=M{x}-N{x}')
    dat.cell(x,16,f'=IF(M{x}=0,0,O{x}/M{x})'); dat.cell(x,17,f'=YEAR(B{x})')
    dat.cell(x,18,f'="Q"&ROUNDUP(MONTH(B{x})/3,0)'); dat.cell(x,19,f'=MONTH(B{x})'); dat.cell(x,20,f'=TEXT(B{x},"mmm")')
    dat.cell(x,21,flag(x))
    for k,sk in enumerate(['Co','Ch','Cu','Ca','Pr','M']): dat.cell(x,22+k,flag(x,sk))
fmt={2:'dd-mmm-yyyy',10:'#,##0.00',11:'0%',12:'#,##0.00',13:'#,##0.00',14:'#,##0.00',15:'#,##0.00',16:'0.0%'}
for x in range(R0,R1+1):
    for j in range(1,NC+1): dat.cell(x,j).font=font(10)
    for j,f in fmt.items(): dat.cell(x,j).number_format=f
for j,w in enumerate([8,13,17,16,16,14,17,22,10,11,10,10,12,12,12,11,7,9,9,8,9,11,11,11,11,11,11],1): dat.column_dimensions[L(j)].width=w
t=Table(displayName='SalesData',ref=f'A{HR}:{L(NC)}{R1}'); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
dat.add_table(t); dat.freeze_panes=f'A{R0}'; dat.sheet_view.showGridLines=False; dat.row_dimensions[HR].height=30
def D(col): return f"Data!${col}${R0}:${col}${R1}"
FL={'all':D('U'),'Co':D('V'),'Ch':D('W'),'Cu':D('X'),'Ca':D('Y'),'Pr':D('Z'),'M':D('AA')}

# filter cells on dashboard
fcells={'sel_Region':'B6','sel_Country':'D6','sel_Channel':'F6','sel_Customer':'H6','sel_Category':'J6','sel_Product':'L6','sel_FromMonth':'N6','sel_ToMonth':'P6'}
for nm,c in fcells.items(): wb.defined_names[nm]=DefinedName(nm,attr_text=f"Dashboard!${c[0]}${c[1:]}")
wb.defined_names['m_From']=DefinedName('m_From',attr_text='Analysis!$C$4'); wb.defined_names['m_To']=DefinedName('m_To',attr_text='Analysis!$C$5')

# ---------- Analysis
A=ana; A.sheet_view.showGridLines=False
A['B1']='Analysis (calculation layer behind the dashboard)'; A['B1'].font=font(14,True,NAVY)
nav(A,1,[(11,12),(13,14),(15,16),(17,18)],'Analysis')
for c,w in zip('ABCDEFGHI',[2,26,14,14,14,14,26,16,16]): A.column_dimensions[c].width=w
def head(r,c,labels):
    for k,l in enumerate(labels):
        x=A.cell(r,c+k,l); x.font=font(10,True,WHITE); x.fill=fill(NAVY); x.alignment=Alignment(horizontal='center',wrap_text=True)
A['B3']='Filter helpers'; A['B3'].font=font(11,True,TEAL)
A['B4']='From month (number)'; A['C4']='=MATCH(sel_FromMonth,Lists!$D$4:$D$15,0)'
A['B5']='To month (number)'; A['C5']='=MATCH(sel_ToMonth,Lists!$D$4:$D$15,0)'
A['B6']='Valid month range?'; A['C6']='=IF(C4<=C5,"OK","From month is after To month - nothing shown")'
A['B8']='KPIs (respect all filters)'; A['B8'].font=font(11,True,TEAL)
K={}
kp=[('Revenue',f'=SUMIFS({D("M")},{FL["all"]},1)','#,##0'),('Gross Profit',f'=SUMIFS({D("O")},{FL["all"]},1)','#,##0'),
    ('Gross Margin %','=IF(C9=0,0,C10/C9)','0.0%'),('Units Sold',f'=SUMIFS({D("I")},{FL["all"]},1)','#,##0'),
    ('Orders',f'=COUNTIFS({FL["all"]},1)','#,##0'),('Avg Order Value','=IF(C13=0,0,C9/C13)','#,##0.00'),
    ('Avg Unit Price (net)','=IF(C12=0,0,C9/C12)','#,##0.00'),('Total Revenue (no filter)',f'=SUM({D("M")})','#,##0'),
    ('Total Gross Profit (no filter)',f'=SUM({D("O")})','#,##0'),('Revenue share of total','=IF(C16=0,0,C9/C16)','0.0%')]
for i,(l,f,nf) in enumerate(kp):
    r=9+i; A.cell(r,2,l).font=font(10); c=A.cell(r,3,f); c.number_format=nf; c.font=font(10); K[l]=f'Analysis!$C${r}'
# monthly (ignores month range; highlights selected months)
A['B21']='Monthly trend (shows all months; months outside the From-To range are greyed)'; A['B21'].font=font(11,True,TEAL)
head(22,2,['Month','Revenue','Gross Profit','Revenue (selected months)','Excluded months'])
for i,m in enumerate(months):
    r=23+i; A.cell(r,2,f'=Lists!$D${4+i}')
    A.cell(r,3,f'=SUMIFS({D("M")},{FL["M"]},1,{D("S")},{i+1})').number_format='#,##0'
    A.cell(r,4,f'=SUMIFS({D("O")},{FL["M"]},1,{D("S")},{i+1})').number_format='#,##0'
    A.cell(r,5,f'=IF(AND({i+1}>=m_From,{i+1}<=m_To),C{r},0)').number_format='#,##0'
    A.cell(r,6,f'=C{r}-E{r}').number_format='#,##0'
    for c in range(2,7): A.cell(r,c).font=font(10)

def block(top,title,items,metric,dimcol,flagkey,selname,nfmt='#,##0',topn=None):
    A.cell(top,2,title).font=font(11,True,TEAL)
    hr=top+1; head(hr,2,['Item','Value (own filter ignored)','Selected','Rank key'])
    n=len(items); f1,f2=hr+1,hr+n
    for i,it in enumerate(items):
        r=hr+1+i
        A.cell(r,2,it).font=font(10)
        c=A.cell(r,3,f'=SUMIFS({D(metric)},{FL[flagkey]},1,{D(dimcol)},B{r})'); c.number_format=nfmt; c.font=font(10)
        d=A.cell(r,4,f'=IF(OR({selname}="All",B{r}={selname}),C{r},0)'); d.number_format=nfmt; d.font=font(10)
        k=A.cell(r,5,f'=C{r}+ROW()/1000000'); k.font=font(9,c=GREY); k.number_format='0.000000'
    head(hr,7,['Sorted item','Selected','Excluded'])
    m=topn or n
    for i in range(m):
        r=hr+1+i
        mt=f'MATCH(LARGE($E${f1}:$E${f2},{i+1}),$E${f1}:$E${f2},0)'
        A.cell(r,7,f'=IF(INDEX($C${f1}:$C${f2},{mt})=0,"",INDEX($B${f1}:$B${f2},{mt}))').font=font(10)
        c=A.cell(r,8,f'=INDEX($D${f1}:$D${f2},{mt})'); c.number_format=nfmt; c.font=font(10)
        c=A.cell(r,9,f'=INDEX($C${f1}:$C${f2},{mt})-INDEX($D${f1}:$D${f2},{mt})'); c.number_format=nfmt; c.font=font(10)
    return dict(hr=hr,f1=f1,f2=f2,s1=hr+1,s2=hr+m,last=hr+max(n,m))
r=37
CAT=block(r,'Revenue by Product Category',cats,'M','G','Ca','sel_Category'); r=CAT['last']+3
CH =block(r,'Revenue by Sales Channel',chans,'M','E','Ch','sel_Channel'); r=CH['last']+3
CU =block(r,'Revenue by Customer Type',custs,'M','F','Cu','sel_Customer'); r=CU['last']+3
CTY=block(r,'Revenue by Country',countries,'M','D','Co','sel_Country'); r=CTY['last']+3
PRD=block(r,'Units Sold by Product (Top 5 shown)',prods,'I','H','Pr','sel_Product',topn=5)
A.sheet_properties.tabColor=TEAL
# highlights
A['G3']='Highlights (respect all filters)'; A['G3'].font=font(11,True,TEAL)
def top_of(b,val_fmt):
    it=f'$B${b["f1"]}:$B${b["f2"]}'; sv=f'$D${b["f1"]}:$D${b["f2"]}'
    return f'=IF($C$9=0,"-",INDEX({it},MATCH(MAX({sv}),{sv},0))&"  ({val_fmt(f"MAX({sv})")})")'
usd=lambda e: f'$"&TEXT({e}/1000,"#,##0.0")&"K'
hl=[('Top country',top_of(CTY,usd)),('Best month','=IF($C$9=0,"-",INDEX($B$23:$B$34,MATCH(MAX($E$23:$E$34),$E$23:$E$34,0))&"  ($"&TEXT(MAX($E$23:$E$34)/1000,"#,##0.0")&"K)")'),
    ('Top category',top_of(CAT,usd)),('Top sales channel',top_of(CH,usd)),('Top customer type',top_of(CU,usd)),
    ('Top product (units)',top_of(PRD,lambda e:f'"&{e}&" units'))]
for i,(l,f) in enumerate(hl):
    A.cell(4+i,7,l).font=font(10); A.cell(4+i,8,f).font=font(10)
HL=[f'Analysis!$H${4+i}' for i in range(6)]

# ---------- Dashboard
S=dash; S.sheet_view.showGridLines=False; S.sheet_view.zoomScale=90
S.column_dimensions['A'].width=2
for c in range(2,20): S.column_dimensions[L(c)].width=10.5
S.column_dimensions['T'].width=2
LAST=73
for r in range(1,LAST+1):
    for c in range(1,21): S.cell(r,c).fill=fill(SAND)
for c in range(2,20):
    for r in (2,3): S.cell(r,c).fill=fill(NAVY)
S.row_dimensions[1].height=8; S.row_dimensions[2].height=26; S.row_dimensions[3].height=18
S['B2']='SALES PERFORMANCE DASHBOARD  |  FY2026'; S['B2'].font=font(18,True,WHITE); S['B2'].alignment=Alignment(vertical='center'); S.merge_cells('B2:K2')
S['B3']='=IF(Analysis!$C$6<>"OK","WARNING: "&Analysis!$C$6,N6&"-"&P6&"  |  Region: "&B6&"  |  Country: "&D6&"  |  Channel: "&F6&"  |  Customer: "&H6&"  |  Category: "&J6&"  |  Product: "&L6)'
S['B3'].font=font(9,False,'CBD5E1'); S['B3'].alignment=Alignment(shrink_to_fit=True); S.merge_cells('B3:K3')
if not IMG: nav(S,2,[(12,13),(14,15),(16,17),(18,19)],'Dashboard')
else:
    S['L2']='Source: 2,707 orders  |  Jan-Dec 2026  |  USD'; S['L2'].font=font(10,False,'CBD5E1'); S['L2'].alignment=Alignment(horizontal='right',vertical='center'); S.merge_cells('L2:S2')
S.row_dimensions[4].height=8; S.row_dimensions[5].height=16; S.row_dimensions[6].height=24
flt=[('B','C','Region',LR,'All'),('D','E','Country',LCo,'All'),('F','G','Sales Channel',LCh,'All'),('H','I','Customer Type',LCu,'All'),
     ('J','K','Category',LCa,'All'),('L','M','Product',LP,'All'),('N','O','From month',LM,'Jan'),('P','Q','To month',LM,'Dec')]
orange=Side(style='thin',color=CORAL)
for a_,b_,l,rng,dflt in flt:
    S[f'{a_}5']=l.upper(); S[f'{a_}5'].font=font(8,True,GREY); S.merge_cells(f'{a_}5:{b_}5')
    S[f'{a_}6']=dflt; S[f'{a_}6'].font=font(10,True,NAVY); S.merge_cells(f'{a_}6:{b_}6')
    S[f'{a_}6'].alignment=Alignment(horizontal='center',vertical='center',shrink_to_fit=True)
    for col in (a_,b_):
        S[f'{col}6'].fill=fill('FFF8E1'); S[f'{col}6'].protection=Protection(locked=False)
        S[f'{col}6'].border=Border(top=orange,bottom=orange,left=orange if col==a_ else None,right=orange if col==b_ else None)
    dv=DataValidation(type='list',formula1='='+rng,allow_blank=False,showErrorMessage=True,errorTitle='Pick from list',error='Choose a value from the dropdown.')
    S.add_data_validation(dv); dv.add(f'{a_}6')
S['R5']='ORDERS SHOWN'; S['R5'].font=font(8,True,GREY); S.merge_cells('R5:S5')
S['R6']=f'={K["Orders"]}'; S['R6'].font=font(14,True,NAVY); S['R6'].number_format='#,##0'; S['R6'].alignment=Alignment(horizontal='center',vertical='center'); S.merge_cells('R6:S6')
for c in ('R6','S6'): S[c].fill=fill(WHITE); S[c].border=Border(top=Side(style='thin',color='D1D5DB'),bottom=Side(style='thin',color='D1D5DB'))
S.row_dimensions[7].height=8
cards=[('TOTAL REVENUE',K['Revenue'],'$#,##0',f'=TEXT({K["Revenue share of total"]},"0.0%")&" of full-year revenue"',TEAL),
       ('GROSS PROFIT',K['Gross Profit'],'$#,##0',f'="Margin "&TEXT({K["Gross Margin %"]},"0.0%")',TEAL),
       ('GROSS MARGIN',K['Gross Margin %'],'0.0%','="Profit as % of revenue"',TEAL),
       ('UNITS SOLD',K['Units Sold'],'#,##0',f'=TEXT({K["Orders"]},"#,##0")&" orders"',CORAL),
       ('AVG ORDER VALUE',K['Avg Order Value'],'$#,##0.00','="Revenue per order"',CORAL),
       ('AVG UNIT PRICE',K['Avg Unit Price (net)'],'$#,##0.00','="Net revenue per unit"',CORAL)]
for r_,h in zip((8,9,10,11),(16,22,22,16)): S.row_dimensions[r_].height=h
for i,(lab,ref,nf,sub,col) in enumerate(cards):
    c1=2+i*3; c2=c1+2
    for r_ in range(8,12):
        for c in range(c1,c2+1): S.cell(r_,c).fill=fill(WHITE)
    for c in range(c1,c2+1): S.cell(8,c).border=Border(top=Side(style='thick',color=col))
    S.cell(8,c1,lab).font=font(8,True,GREY); S.merge_cells(start_row=8,start_column=c1,end_row=8,end_column=c2)
    v=S.cell(9,c1,'='+ref); v.font=font(22,True,NAVY); v.number_format=nf; v.alignment=Alignment(horizontal='left',vertical='center',indent=1)
    S.merge_cells(start_row=9,start_column=c1,end_row=10,end_column=c2)
    s_=S.cell(11,c1,sub); s_.font=font(8,False,GREY); S.merge_cells(start_row=11,start_column=c1,end_row=11,end_column=c2)
    S.cell(8,c1).alignment=Alignment(indent=1); S.cell(11,c1).alignment=Alignment(indent=1)
S.row_dimensions[12].height=8; S.row_dimensions[13].height=15; S.row_dimensions[14].height=20
for i,(lab,ref) in enumerate(zip(['TOP COUNTRY','BEST MONTH','TOP CATEGORY','TOP CHANNEL','TOP CUSTOMER TYPE','TOP PRODUCT'],HL)):
    c1=2+i*3; c2=c1+2
    for r_ in (13,14):
        for c in range(c1,c2+1): S.cell(r_,c).fill=fill(LIGHT)
    S.cell(13,c1,lab).font=font(7,True,TEAL); S.cell(13,c1).alignment=Alignment(indent=1,vertical='bottom'); S.merge_cells(start_row=13,start_column=c1,end_row=13,end_column=c2)
    v=S.cell(14,c1,'='+ref); v.font=font(10,True,NAVY); v.alignment=Alignment(indent=1,vertical='center'); S.merge_cells(start_row=14,start_column=c1,end_row=14,end_column=c2)
S.row_dimensions[15].height=8
def panel(r1,r2,c1,c2,title):
    for r_ in range(r1,r2+1):
        for c in range(c1,c2+1): S.cell(r_,c).fill=fill(WHITE)
    t=S.cell(r1,c1,title); t.font=font(11,True,NAVY); t.alignment=Alignment(indent=1,vertical='center')
    S.merge_cells(start_row=r1,start_column=c1,end_row=r1,end_column=c2); S.row_dimensions[r1].height=22
    for c in range(c1,c2+1): S.cell(r1,c).border=Border(bottom=Side(style='medium',color=TEAL if c1==2 else CORAL))
panel(16,33,2,11,'Monthly Revenue & Gross Profit'); panel(16,33,12,19,'Revenue by Product Category')
S.row_dimensions[34].height=8
panel(35,52,2,11,'Revenue Map of Europe (tile map)'); panel(35,52,12,19,'Top 5 Bestselling Products (units)')
S.row_dimensions[53].height=8
panel(54,72,2,11,'Top 10 Countries by Revenue'); panel(54,72,12,19,'Revenue Mix')

def axfont(ax,sz=800):
    cp=CharacterProperties(sz=sz,solidFill=GREY,latin=DFont(typeface='Arial'))
    ax.txPr=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp),endParaRPr=cp)])
def polish(ch,kf=False):
    ch.roundedCorners=False; ch.y_axis.scaling.min=0
    ch.graphical_properties=GraphicalProperties(ln=LineProperties(noFill=True))
    for ax in (ch.x_axis,ch.y_axis): ax.delete=False; axfont(ax)
    ch.y_axis.majorGridlines.spPr=GraphicalProperties(ln=LineProperties(solidFill='E5E7EB'))
    ch.y_axis.spPr=GraphicalProperties(ln=LineProperties(noFill=True)); ch.x_axis.spPr=GraphicalProperties(ln=LineProperties(solidFill='D1D5DB'))
    if kf: ch.y_axis.numFmt='$#,##0,"K"'; ch.y_axis.number_format='$#,##0,"K"'
def lbl(ser,fmt):
    ser.dLbls=DataLabelList(); ser.dLbls.showVal=True; ser.dLbls.showSerName=False; ser.dLbls.showCatName=False; ser.dLbls.showLegendKey=False; ser.dLbls.numFmt=fmt
    cp=CharacterProperties(sz=800,b=True,solidFill='FFFFFF',latin=DFont(typeface='Arial'))
    ser.dLbls.txPr=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp),endParaRPr=cp)])
OTHER='D1D5DB'
def sbar(blk,anchor,w,h,color,kf=True,horiz=True,lblfmt='$#,##0,"K";;;'):
    c=BarChart(); c.type='bar' if horiz else 'col'; c.grouping='stacked'; c.overlap=100; c.gapWidth=45
    c.add_data(Reference(A,min_col=8,max_col=9,min_row=blk['hr'],max_row=blk['s2']),titles_from_data=True)
    c.set_categories(Reference(A,min_col=7,min_row=blk['s1'],max_row=blk['s2']))
    c.series[0].graphicalProperties=GraphicalProperties(solidFill=color); c.series[1].graphicalProperties=GraphicalProperties(solidFill=OTHER)
    c.x_axis.scaling.orientation='maxMin'; c.legend=None; polish(c,kf); lbl(c.series[0],lblfmt)
    c.width=w; c.height=h; S.add_chart(c,anchor); return c
WL,WR=20.4,16.3
bc=BarChart(); bc.type='col'; bc.grouping='stacked'; bc.overlap=100; bc.gapWidth=55
bc.add_data(Reference(A,min_col=5,max_col=6,min_row=22,max_row=34),titles_from_data=True)
bc.set_categories(Reference(A,min_col=2,min_row=23,max_row=34))
bc.series[0].graphicalProperties=GraphicalProperties(solidFill=TEAL); bc.series[1].graphicalProperties=GraphicalProperties(solidFill=OTHER)
ln=LineChart(); ln.add_data(Reference(A,min_col=4,min_row=22,max_row=34),titles_from_data=True)
ln.series[0].graphicalProperties.line.solidFill=CORAL; ln.series[0].graphicalProperties.line.width=28000
ln.series[0].marker.symbol='circle'; ln.series[0].marker.size=6; ln.series[0].marker.graphicalProperties=GraphicalProperties(solidFill=CORAL); ln.series[0].smooth=False
bc.legend.position='t'; polish(bc,True); bc+=ln; bc.width=WL; bc.height=8.7; S.add_chart(bc,'B17')
sbar(CAT,'L17',WR,8.7,NAVY)
sbar(PRD,'L36',WR,11.0,CORAL,kf=False,lblfmt='#,##0;;;')
sbar(dict(CTY,s2=CTY['hr']+10),'B55',WL,8.7,TEAL)
c_a=sbar(CH,'L56',8.0,8.4,TEAL); c_b=sbar(CU,'P56',8.0,8.4,CORAL)
if IMG:
    for c_ in (c_a,c_b): c_.y_axis.delete=True; c_.y_axis.majorGridlines=None
S['L55']='BY SALES CHANNEL'; S['P55']='BY CUSTOMER TYPE'
for a_ in ('L55','P55'): S[a_].font=font(8,True,GREY); S[a_].alignment=Alignment(indent=1)
# tile map
tiles={'Norway':(0,1),'Sweden':(0,2),'Finland':(0,3),'Ireland':(1,0),'United Kingdom':(1,1),'Denmark':(1,2),
       'Netherlands':(2,1),'Germany':(2,2),'Poland':(2,3),'Belgium':(3,0),'France':(3,1),'Switzerland':(3,2),'Austria':(3,3),'Czechia':(3,4),
       'Portugal':(4,0),'Spain':(4,1),'Italy':(4,2)}
assert set(tiles)==set(countries)
S.row_dimensions[36].height=6
for r_ in range(37,47): S.row_dimensions[r_].height=22
valcells=[]
for nme,(tr,tc) in tiles.items():
    r_=37+tr*2; c=2+tc*2
    S.cell(r_,c,nme).font=font(8,True,WHITE); S.cell(r_,c).alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    S.merge_cells(start_row=r_,start_column=c,end_row=r_,end_column=c+1)
    S.cell(r_+1,c,f'=INDEX(Analysis!$C${CTY["f1"]}:$C${CTY["f2"]},MATCH("{nme}",Analysis!$B${CTY["f1"]}:$B${CTY["f2"]},0))')
    S.cell(r_+1,c).number_format='$#,##0'; S.cell(r_+1,c).font=font(9,True,WHITE); S.cell(r_+1,c).alignment=Alignment(horizontal='center',vertical='center')
    S.merge_cells(start_row=r_+1,start_column=c,end_row=r_+1,end_column=c+1)
    for cc in (c,c+1): S.cell(r_,cc).fill=fill(NAVY)
    valcells.append(S.cell(r_+1,c).coordinate)
    S.conditional_formatting.add(f'{L(c)}{r_}:{L(c+1)}{r_}',FormulaRule(formula=[f'${L(c)}${r_}=sel_Country'],fill=PatternFill('solid',bgColor=CORAL,fgColor=CORAL),stopIfTrue=False))
S.conditional_formatting.add(' '.join(f'{v}:{L(openpyxl.utils.cell.column_index_from_string(v[0])+1)}{v[1:]}' for v in valcells),
    ColorScaleRule(start_type='min',start_color='9FD3D4',mid_type='percentile',mid_value=50,mid_color='2BA3A5',end_type='max',end_color='0B4F52'))
S['B48']='Tiles sit roughly where each country is on the map. Shade = revenue (light = low, dark = high). The selected country gets an orange label. With the macro installed, click a tile to filter the whole dashboard.'
S['B48'].font=font(8,False,GREY,True); S['B48'].alignment=Alignment(wrap_text=True,vertical='top',indent=1); S.merge_cells('B48:K51')
if IMG: S['B48']='Tiles sit roughly where each country is on the map. Shade = revenue (light = low, dark = high).'
S.freeze_panes='A7'; S.sheet_properties.tabColor=NAVY
S.page_setup.orientation='landscape'; S.page_setup.paperSize=9
S.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); S.page_setup.fitToWidth=1; S.page_setup.fitToHeight=1
S.print_area=f'A1:T{LAST}'; S.print_options.horizontalCentered=True
S.page_margins.left=S.page_margins.right=0.3; S.page_margins.top=S.page_margins.bottom=0.3
S.protection.sheet=True; S.protection.selectLockedCells=False; S.protection.formatCells=True
wb.properties.title='Sales Performance Dashboard FY2026'; wb.properties.creator='Sales Analytics'

# ---------- Guide
G=gd; G.sheet_view.showGridLines=False
G.column_dimensions['A'].width=2; G.column_dimensions['B'].width=26; G.column_dimensions['C'].width=95
G['B1']='Guide'; G['B1'].font=font(16,True,NAVY)
nav(G,1,[(5,5),(6,6),(7,7),(8,8)],'Guide')
for c in 'EFGH': G.column_dimensions[c].width=13
guide=[('How to use','Use the 8 orange dropdowns on the Dashboard (Region, Country, Channel, Customer, Category, Product, From/To month). Everything recalculates instantly. Set a dropdown back to "All" to clear it.'),
('Cross-filtering','Every chart ignores its OWN filter and shows the selected bar in colour while other bars turn grey (Power BI style); all other visuals, KPIs and highlights filter to the selection. Without macros you choose the item in the dropdown; with the macro in CrossFilter_Macro.txt you click the bar, month or map tile directly. Clicking it again clears it. Clicking the ORDERS SHOWN box resets all filters.'),
('Navigation','Use the Dashboard | Analysis | Data | Guide buttons at the top of every sheet.'),
('Data sheet','Table "SalesData" - 2,707 orders. Revenue, COGS, Gross Profit, Margin, Year, Quarter and Month are formulas. Column U "In Filter" = 1 when a row passes every filter; columns V:AA are the same test with one filter left out (used so a chart can still show the bars that are not selected).'),
('To add data','Paste new rows under the table, extend the formulas in M:AA, and widen the ranges (rows 4 to %d).'%R1),
('Analysis sheet','All calculations: KPIs, monthly trend and sorted blocks for category, channel, customer type, country and top-5 products. Highlights strip is in G3:H9.'),
('Definitions','Revenue = Units x Unit Price x (1 - Discount). COGS = Units x Unit Cost. Gross Profit = Revenue - COGS. Avg Order Value = Revenue / Orders. Avg Unit Price = Revenue / Units.'),
('Protection','The Dashboard sheet is protected (no password); only orange filter cells are editable. To edit the layout: Review > Unprotect Sheet. The macro re-protects it in a chart-friendly way each time the file opens.'),
('Limits','Dropdowns/macro clicks replace slicers and the timeline; cells replace pivot tables and the filled map. Chart-click filtering needs the macro and desktop Excel (Windows or Mac); it does not run in Excel for the web or mobile.')]
for i,(a_,b_) in enumerate(guide):
    r_=3+i*2; G.cell(r_,2,a_).font=font(10,True,TEAL); c=G.cell(r_,3,b_); c.font=font(10); c.alignment=Alignment(wrap_text=True,vertical='top'); G.row_dimensions[r_].height=52
    G.cell(r_,2).alignment=Alignment(vertical='top')
G.sheet_properties.tabColor='9CA3AF'
wb.calculation.fullCalcOnLoad=True
wb.save(OUT)
print(f'Saved {OUT} ({N} orders)')

src=OUT; tmp=src+'.tmp'
with zipfile.ZipFile(src) as zi, zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        data=zi.read(it.filename)
        if it.filename.startswith('xl/charts/chart') and it.filename.endswith('.xml'):
            t=data.decode('utf8'); t=re.sub(r'<numFmt formatCode="([^"]*)"\s*/>',r'<numFmt formatCode="\1" sourceLinked="0"/>',t); data=t.encode('utf8')
        zo.writestr(it,data)
shutil.move(tmp,src); 
